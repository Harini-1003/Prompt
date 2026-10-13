from pathlib import Path
from typing import Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_to_pdf(paper_data: Dict[str, Any], output_path: Path, include_answers: bool = True) -> Path:
    """Export formatted university-style exam question paper and answer key as PDF."""
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        alignment=1, # Center
        spaceAfter=4
    )
    sub_title_style = ParagraphStyle(
        'SubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        alignment=1,
        spaceAfter=8
    )
    section_title = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        spaceBefore=12,
        spaceAfter=6,
        textColor=colors.HexColor("#1e3a8a")
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12
    )

    story = []

    # Header / College Title
    college = paper_data.get("college_name", "COLLEGE OF ENGINEERING & TECHNOLOGY")
    exam = paper_data.get("exam_name", "END SEMESTER DEGREE EXAMINATION")
    subject = paper_data.get("subject_name", "Computer Science")
    code = paper_data.get("subject_code", "CS101")
    duration = paper_data.get("duration_minutes", 180)
    total_marks = paper_data.get("total_marks", 100)

    story.append(Paragraph(college.upper(), title_style))
    story.append(Paragraph(f"{exam.upper()}", sub_title_style))
    story.append(Spacer(1, 4))

    # Meta table
    meta_data = [
        [
            Paragraph(f"<b>Subject:</b> {subject} ({code})", cell_style),
            Paragraph(f"<b>Time:</b> {duration} Minutes", cell_style),
            Paragraph(f"<b>Max Marks:</b> {total_marks}", cell_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[240, 150, 150])
    meta_table.setStyle(TableStyle([
        ('LINEBELOW', (0,0), (-1,-1), 1.5, colors.HexColor("#1e3a8a")),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # General Instructions
    instructions_text = "<b>Instructions:</b> Answer all questions. Clearly mention Question Numbers. Marks and CO mappings are indicated in the table."
    story.append(Paragraph(instructions_text, cell_style))
    story.append(Spacer(1, 10))

    # Sections & Questions
    sections = paper_data.get("sections", [])
    for sec in sections:
        sec_name = sec.get("section_name", "Part")
        sec_type = sec.get("section_type", "")
        story.append(Paragraph(f"<b>{sec_name}</b> - {sec_type}", section_title))

        # Question table headers
        table_rows = [
            [
                Paragraph("<b>Q#</b>", cell_bold),
                Paragraph("<b>Question</b>", cell_bold),
                Paragraph("<b>CO</b>", cell_bold),
                Paragraph("<b>Bloom Level</b>", cell_bold),
                Paragraph("<b>Marks</b>", cell_bold)
            ]
        ]

        for q in sec.get("questions", []):
            table_rows.append([
                Paragraph(str(q.get("question_number", "")), cell_style),
                Paragraph(str(q.get("question_text", "")), cell_style),
                Paragraph(str(q.get("course_outcome", "")), cell_style),
                Paragraph(str(q.get("bloom_level", "")), cell_style),
                Paragraph(str(q.get("marks", "")), cell_style)
            ])

        q_table = Table(table_rows, colWidths=[40, 310, 50, 90, 50])
        q_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#f1f5f9")),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(q_table)
        story.append(Spacer(1, 10))

    # Optional Answer Key Section
    if include_answers:
        story.append(PageBreak())
        story.append(Paragraph("CONFIDENTIAL - OFFICIAL ANSWER KEY & MARKING SCHEME", title_style))
        story.append(Paragraph(f"{subject} ({code}) - {exam}", sub_title_style))
        story.append(Spacer(1, 8))

        all_q = paper_data.get("all_questions", [])
        ans_rows = [
            [
                Paragraph("<b>Q#</b>", cell_bold),
                Paragraph("<b>Expected Model Answer</b>", cell_bold),
                Paragraph("<b>Marking Scheme Hints</b>", cell_bold),
                Paragraph("<b>Marks</b>", cell_bold)
            ]
        ]
        for q in all_q:
            ans_rows.append([
                Paragraph(str(q.get("question_number", "")), cell_style),
                Paragraph(str(q.get("model_answer", "N/A")), cell_style),
                Paragraph(str(q.get("marking_scheme", "N/A")), cell_style),
                Paragraph(str(q.get("marks", "")), cell_style)
            ])

        ans_table = Table(ans_rows, colWidths=[40, 280, 170, 50])
        ans_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#e2e8f0")),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#94a3b8")),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(ans_table)

    doc.build(story)
    return output_path

def export_to_docx(paper_data: Dict[str, Any], output_path: Path, include_answers: bool = True) -> Path:
    """Export formatted university-style question paper and answer key as DOCX."""
    doc = docx.Document()

    college = paper_data.get("college_name", "COLLEGE OF ENGINEERING & TECHNOLOGY")
    exam = paper_data.get("exam_name", "END SEMESTER DEGREE EXAMINATION")
    subject = paper_data.get("subject_name", "Computer Science")
    code = paper_data.get("subject_code", "CS101")
    duration = paper_data.get("duration_minutes", 180)
    total_marks = paper_data.get("total_marks", 100)

    # Header
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = h1.add_run(college.upper() + "\n")
    run1.bold = True
    run1.font.size = Pt(14)
    run2 = h1.add_run(exam.upper())
    run2.bold = True
    run2.font.size = Pt(11)

    # Details paragraph
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Subject: {subject} ({code})       Time: {duration} mins       Max Marks: {total_marks}\n").bold = True
    p_meta.add_run("Instructions: Answer all questions. Mappings and Bloom levels are specified.")

    sections = paper_data.get("sections", [])
    for sec in sections:
        sec_h = doc.add_heading(f"{sec.get('section_name')} - {sec.get('section_type')}", level=2)
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Q#"
        hdr_cells[1].text = "Question"
        hdr_cells[2].text = "CO"
        hdr_cells[3].text = "Bloom's Level"
        hdr_cells[4].text = "Marks"
        for c in hdr_cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True

        for q in sec.get("questions", []):
            row_cells = table.add_row().cells
            row_cells[0].text = str(q.get("question_number", ""))
            row_cells[1].text = str(q.get("question_text", ""))
            row_cells[2].text = str(q.get("course_outcome", ""))
            row_cells[3].text = str(q.get("bloom_level", ""))
            row_cells[4].text = str(q.get("marks", ""))

    if include_answers:
        doc.add_page_break()
        ans_h = doc.add_heading("CONFIDENTIAL - ANSWER KEY & MARKING SCHEME", level=1)
        ans_table = doc.add_table(rows=1, cols=4)
        ans_table.style = 'Table Grid'
        hdr = ans_table.rows[0].cells
        hdr[0].text = "Q#"
        hdr[1].text = "Expected Model Answer"
        hdr[2].text = "Marking Scheme"
        hdr[3].text = "Marks"
        for c in hdr:
            for p in c.paragraphs:
                for r in p.runs:
                    r.bold = True

        for q in paper_data.get("all_questions", []):
            row = ans_table.add_row().cells
            row[0].text = str(q.get("question_number", ""))
            row[1].text = str(q.get("model_answer", "N/A"))
            row[2].text = str(q.get("marking_scheme", "N/A"))
            row[3].text = str(q.get("marks", ""))

    doc.save(str(output_path))
    return output_path
